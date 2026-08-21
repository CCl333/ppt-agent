from __future__ import annotations

import base64
import io
import re
from xml.etree import ElementTree as ET

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.connector import Connector
from pptx.util import Emu, Pt

from app.services.font_policy import normalize_family
from app.services.svg import CANVAS_HEIGHT, CANVAS_WIDTH, _ensure_svg_xmlns, _local_tag, _parse_svg
from app.services.svg_contract import estimated_text_width, flatten_group_translates, parse_length

EMU_PER_PX = 9525
SLIDE_WIDTH_EMU = 12192000
SLIDE_HEIGHT_EMU = 6858000
PT_PER_PX = 0.75
ASCENT_RATIO = 0.88
_DATA_URI_RE = re.compile(r"^data:image/([a-zA-Z0-9+.-]+);base64,(.+)$", re.DOTALL)
MAX_IMAGE_BYTES = 8 * 1024 * 1024
_PATH_TOKEN_RE = re.compile(r"[A-Za-z]|[-+]?\d*\.?\d+(?:e[-+]?\d+)?")


def emu(px: float) -> int:
    return round(float(px) * EMU_PER_PX)


def font_sz(font_px: float) -> int:
    return round(float(font_px) * PT_PER_PX * 100)


def add_svg_slide(presentation: Presentation, svg_markup: str) -> None:
    blank = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(blank)
    try:
        slide.shapes.turbo_add_enabled = True
    except Exception:
        pass
    root = _parse_svg(_ensure_svg_xmlns(svg_markup))
    flatten_group_translates(root)
    for elem in list(root):
        _emit_element(slide, elem)


def _emit_element(slide, elem: ET.Element) -> None:
    tag = _local_tag(elem)
    if tag in {"title", "desc"}:
        return
    if tag == "g":
        for child in list(elem):
            _emit_element(slide, child)
        return
    if _is_hidden(elem):
        return
    emitters = {
        "rect": _emit_rect,
        "circle": lambda slide, node: _emit_ellipse(slide, node, circle=True),
        "ellipse": lambda slide, node: _emit_ellipse(slide, node, circle=False),
        "line": _emit_line,
        "polygon": lambda slide, node: _emit_poly(slide, node, closed=True),
        "polyline": lambda slide, node: _emit_poly(slide, node, closed=False),
        "path": _emit_path,
        "text": _emit_text,
        "image": _emit_image,
    }
    emitter = emitters.get(tag)
    if emitter is None:
        raise RuntimeError(f"导出器无法翻译图元 <{tag}>，禁止静默跳过")
    emitter(slide, elem)


def _emit_rect(slide, elem: ET.Element) -> None:
    x = parse_length(elem.get("x"))
    y = parse_length(elem.get("y"))
    width = parse_length(elem.get("width"))
    height = parse_length(elem.get("height"))
    rx = parse_length(elem.get("rx") or elem.get("ry"))
    if width <= 0 or height <= 0:
        return
    preset = MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0 else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(preset, Emu(emu(x)), Emu(emu(y)), Emu(emu(width)), Emu(emu(height)))
    if rx > 0:
        half = min(width, height) / 2
        adj = int(round((rx / half) * 100000)) if half else 0
        _set_round_rect_adj(shape, max(0, min(adj, 50000)))
    _apply_paint(shape, elem)


def _emit_ellipse(slide, elem: ET.Element, *, circle: bool) -> None:
    if circle:
        r = parse_length(elem.get("r"))
        cx = parse_length(elem.get("cx"))
        cy = parse_length(elem.get("cy"))
        x, y, w, h = cx - r, cy - r, r * 2, r * 2
    else:
        rx = parse_length(elem.get("rx"))
        ry = parse_length(elem.get("ry"))
        cx = parse_length(elem.get("cx"))
        cy = parse_length(elem.get("cy"))
        x, y, w, h = cx - rx, cy - ry, rx * 2, ry * 2
    if w <= 0 or h <= 0:
        return
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(h)))
    _apply_paint(shape, elem)


def _emit_line(slide, elem: ET.Element) -> None:
    x1 = parse_length(elem.get("x1"))
    y1 = parse_length(elem.get("y1"))
    x2 = parse_length(elem.get("x2"))
    y2 = parse_length(elem.get("y2"))
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Emu(emu(x1)),
        Emu(emu(y1)),
        Emu(emu(x2)),
        Emu(emu(y2)),
    )
    _apply_paint(shape, elem, line_fallback=True)


def _emit_poly(slide, elem: ET.Element, *, closed: bool) -> None:
    nums = [float(item) for item in re.findall(r"[-+]?\d*\.?\d+(?:e[-+]?\d+)?", elem.get("points") or "")]
    points = [(nums[index], nums[index + 1]) for index in range(0, len(nums) - 1, 2)]
    _emit_points(slide, elem, points, closed=closed)


def _emit_path(slide, elem: ET.Element) -> None:
    contours = parse_path_contours(elem.get("d") or "")
    if not contours:
        return
    first_points, first_closed = contours[0]
    if len(first_points) < 2:
        return
    builder = slide.shapes.build_freeform(Emu(emu(first_points[0][0])), Emu(emu(first_points[0][1])))
    if len(first_points) > 1:
        builder.add_line_segments(
            [(Emu(emu(x)), Emu(emu(y))) for x, y in first_points[1:]],
            close=first_closed,
        )
    for points, closed in contours[1:]:
        if len(points) < 2:
            continue
        builder.move_to(Emu(emu(points[0][0])), Emu(emu(points[0][1])))
        builder.add_line_segments(
            [(Emu(emu(x)), Emu(emu(y))) for x, y in points[1:]],
            close=closed,
        )
    shape = builder.convert_to_shape()
    _apply_paint(shape, elem)


def _emit_points(slide, elem: ET.Element, points: list[tuple[float, float]], *, closed: bool) -> None:
    if len(points) < 2:
        return
    builder = slide.shapes.build_freeform(Emu(emu(points[0][0])), Emu(emu(points[0][1])))
    builder.add_line_segments(
        [(Emu(emu(x)), Emu(emu(y))) for x, y in points[1:]],
        close=closed,
    )
    shape = builder.convert_to_shape()
    _apply_paint(shape, elem)


def _emit_text(slide, elem: ET.Element) -> None:
    content = _text_content(elem).strip()
    if not content:
        return
    font_px = parse_length(elem.get("font-size"), 16.0)
    x = parse_length(elem.get("x"))
    y = parse_length(elem.get("y"))
    width = max(estimated_text_width(content, font_px) + font_px * 0.4, font_px)
    height = font_px * 1.35
    top = y - font_px * ASCENT_RATIO
    anchor = (elem.get("text-anchor") or "start").strip()
    if anchor == "end":
        left = x - width
        align = PP_ALIGN.RIGHT
    elif anchor in {"middle", "center"}:
        left = x - width / 2
        align = PP_ALIGN.CENTER
    else:
        left = x
        align = PP_ALIGN.LEFT
    shape = slide.shapes.add_textbox(Emu(emu(left)), Emu(emu(top)), Emu(emu(width)), Emu(emu(height)))
    _configure_textbox(shape, align)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.alignment = align
    if paragraph.runs:
        run = paragraph.runs[0]
    else:
        run = paragraph.add_run()
    run.text = content
    _apply_run_font(run, elem, font_px)
    fill, opacity = _fill_color(elem)
    if fill:
        run.font.color.rgb = _rgb(fill)
        if opacity < 0.999:
            _set_run_alpha(run, opacity)


def _emit_image(slide, elem: ET.Element) -> None:
    href = elem.get("href") or elem.get("{http://www.w3.org/1999/xlink}href") or ""
    payload = _decode_image_href(href)
    x = parse_length(elem.get("x"))
    y = parse_length(elem.get("y"))
    width = parse_length(elem.get("width"), CANVAS_WIDTH)
    height = parse_length(elem.get("height"), CANVAS_HEIGHT)
    if width <= 0 or height <= 0:
        return
    slide.shapes.add_picture(io.BytesIO(payload), Emu(emu(x)), Emu(emu(y)), Emu(emu(width)), Emu(emu(height)))


def _configure_textbox(shape, align: PP_ALIGN) -> None:
    frame = shape.text_frame
    frame.word_wrap = False
    nv_sp = shape._element.find(qn("p:nvSpPr"))
    if nv_sp is not None:
        c_nv_sp = nv_sp.find(qn("p:cNvSpPr"))
        if c_nv_sp is not None:
            c_nv_sp.set("txBox", "1")
    body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if body_pr is None:
        return
    body_pr.set("wrap", "none")
    body_pr.set("lIns", "0")
    body_pr.set("tIns", "0")
    body_pr.set("rIns", "0")
    body_pr.set("bIns", "0")
    for child in list(body_pr):
        local = child.tag.rsplit("}", 1)[-1]
        if local in {"spAutoFit", "noAutofit", "normAutofit"}:
            body_pr.remove(child)
    etree.SubElement(body_pr, qn("a:spAutoFit"))


def _apply_run_font(run, elem: ET.Element, font_px: float) -> None:
    run.font.size = Pt(font_px * PT_PER_PX)
    weight = (elem.get("font-weight") or "").strip().lower()
    run.font.bold = weight in {"bold", "700", "800", "900"} or (weight.isdigit() and int(weight) >= 600)
    family = _primary_font(elem.get("font-family") or "")
    run.font.name = family
    rpr = run._r.find(qn("a:rPr"))
    if rpr is None:
        return
    rpr.set("sz", str(font_sz(font_px)))
    for tag_name in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag_name))
        if node is None:
            node = etree.SubElement(rpr, qn(tag_name))
        node.set("typeface", family)


def _apply_paint(shape, elem: ET.Element, *, line_fallback: bool = False) -> None:
    fill, fill_opacity = _fill_color(elem)
    stroke, stroke_opacity = _stroke_color(elem)
    stroke_width = parse_length(elem.get("stroke-width"), 1.0 if stroke or line_fallback else 0.0)
    if line_fallback and not stroke:
        stroke, stroke_opacity = fill, fill_opacity
        fill = None
    if not isinstance(shape, Connector):
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
            if fill_opacity < 0.999:
                _set_fill_alpha(shape, fill_opacity)
        else:
            shape.fill.background()
    if stroke and stroke_width > 0:
        shape.line.color.rgb = _rgb(stroke)
        shape.line.width = Emu(emu(stroke_width))
        if stroke_opacity < 0.999:
            _set_line_alpha(shape, stroke_opacity)
    elif not line_fallback and not isinstance(shape, Connector):
        shape.line.fill.background()


def _fill_color(elem: ET.Element) -> tuple[str | None, float]:
    fill = (elem.get("fill") or "").strip()
    if fill.lower() in {"", "none", "transparent"}:
        return None, 1.0
    opacity = _combined_opacity(elem, "fill-opacity")
    return _norm_hex(fill), opacity


def _stroke_color(elem: ET.Element) -> tuple[str | None, float]:
    stroke = (elem.get("stroke") or "").strip()
    if stroke.lower() in {"", "none", "transparent"}:
        return None, 1.0
    opacity = _combined_opacity(elem, "stroke-opacity")
    return _norm_hex(stroke), opacity


def _combined_opacity(elem: ET.Element, paint_attr: str) -> float:
    base = parse_length(elem.get("opacity"), 1.0)
    paint = parse_length(elem.get(paint_attr), 1.0)
    return max(0.0, min(base * paint, 1.0))


def _set_round_rect_adj(shape, adj: int) -> None:
    prst = shape._element.spPr.find(qn("a:prstGeom"))
    if prst is None:
        return
    av_lst = prst.find(qn("a:avLst"))
    if av_lst is None:
        av_lst = etree.SubElement(prst, qn("a:avLst"))
    else:
        for child in list(av_lst):
            av_lst.remove(child)
    gd = etree.SubElement(av_lst, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {adj}")


def _set_fill_alpha(shape, opacity: float) -> None:
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    _set_alpha(srgb, opacity)


def _set_line_alpha(shape, opacity: float) -> None:
    ln = shape._element.spPr.find(qn("a:ln"))
    if ln is None:
        return
    solid = ln.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    _set_alpha(srgb, opacity)


def _set_run_alpha(run, opacity: float) -> None:
    rpr = run._r.find(qn("a:rPr"))
    if rpr is None:
        return
    solid = rpr.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    _set_alpha(srgb, opacity)


def _set_alpha(srgb, opacity: float) -> None:
    for child in list(srgb):
        if child.tag.rsplit("}", 1)[-1] == "alpha":
            srgb.remove(child)
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", str(int(round(opacity * 100000))))


def _rgb(color: str) -> RGBColor:
    raw = _norm_hex(color)[1:]
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _norm_hex(value: str) -> str:
    text = value.strip()
    if text.startswith("#") and len(text) == 4:
        return "#" + "".join(ch * 2 for ch in text[1:]).upper()
    if text.startswith("#") and len(text) >= 7:
        return text[:7].upper()
    raise RuntimeError(f"无法解析颜色: {value}")


def _primary_font(family: str) -> str:
    return normalize_family(family) or "Microsoft YaHei"


def _text_content(elem: ET.Element) -> str:
    parts: list[str] = []
    if elem.text:
        parts.append(elem.text)
    for child in list(elem):
        tag = _local_tag(child)
        if tag == "tspan":
            for attr in ("x", "y", "dx", "dy"):
                if (child.get(attr) or "").strip():
                    raise RuntimeError("tspan 不得设置 x/y/dx/dy，导出器无法保留分片坐标")
            parts.append("".join(child.itertext()))
        elif tag not in {"title", "desc"}:
            raise RuntimeError(f"text 内含有无法翻译的子节点 <{tag}>")
        if child.tail:
            parts.append(child.tail)
    return "".join(parts)


def _is_hidden(elem: ET.Element) -> bool:
    display = (elem.get("display") or "").strip().lower()
    visibility = (elem.get("visibility") or "").strip().lower()
    opacity = parse_length(elem.get("opacity"), 1.0)
    return display == "none" or visibility == "hidden" or opacity <= 0


def _decode_image_href(href: str) -> bytes:
    match = _DATA_URI_RE.match((href or "").strip())
    if not match:
        raise RuntimeError("image 只允许 data:image/...;base64 URI，禁止文件路径或网络地址")
    raw = re.sub(r"\s+", "", match.group(2))
    if len(raw) > MAX_IMAGE_BYTES * 4 // 3 + 8:
        raise RuntimeError("image data URI 超过 8MB 上限")
    try:
        payload = base64.b64decode(raw)
    except Exception as exc:
        raise RuntimeError("image data URI 无法解码") from exc
    if not payload:
        raise RuntimeError("image data URI 为空")
    if len(payload) > MAX_IMAGE_BYTES:
        raise RuntimeError("image data URI 超过 8MB 上限")
    return payload


def parse_path_contours(path_d: str) -> list[tuple[list[tuple[float, float]], bool]]:
    tokens = _PATH_TOKEN_RE.findall(path_d or "")
    contours: list[tuple[list[tuple[float, float]], bool]] = []
    points: list[tuple[float, float]] = []
    closed = False
    command = ""
    cx = cy = sx = sy = 0.0
    args: list[float] = []

    def flush() -> None:
        nonlocal points, closed, sx, sy
        if len(points) >= 2:
            contours.append((points, closed))
        points = []
        closed = False

    def take(values: list[float], count: int) -> list[list[float]]:
        groups: list[list[float]] = []
        for index in range(0, len(values) - count + 1, count):
            groups.append(values[index : index + count])
        return groups

    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.isalpha():
            command = token
            index += 1
            args = []
            while index < len(tokens) and not tokens[index].isalpha():
                args.append(float(tokens[index]))
                index += 1
            if command in {"M", "m"}:
                groups = take(args, 2)
                if not groups:
                    continue
                if points:
                    flush()
                x, y = groups[0]
                if command == "m":
                    x, y = cx + x, cy + y
                cx, cy = x, y
                sx, sy = x, y
                points = [(x, y)]
                for x, y in groups[1:]:
                    if command == "m":
                        x, y = cx + x, cy + y
                    cx, cy = x, y
                    points.append((x, y))
            elif command in {"L", "l"}:
                for x, y in take(args, 2):
                    if command == "l":
                        x, y = cx + x, cy + y
                    cx, cy = x, y
                    points.append((x, y))
            elif command in {"H", "h"}:
                for x in args:
                    if command == "h":
                        x = cx + x
                    cx = x
                    points.append((cx, cy))
            elif command in {"V", "v"}:
                for y in args:
                    if command == "v":
                        y = cy + y
                    cy = y
                    points.append((cx, cy))
            elif command in {"Z", "z"}:
                closed = True
                cx, cy = sx, sy
                flush()
            else:
                raise RuntimeError(f"path 含有无法翻译的指令: {command}")
        else:
            index += 1
    if points:
        flush()
    return contours
