from __future__ import annotations

from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from app.services.svg import extract_and_validate_svg
from app.services.svg_pptx import SLIDE_HEIGHT_EMU, SLIDE_WIDTH_EMU, add_svg_slide


def preflight_shapes_slide(svg_markup: str) -> dict[str, Any]:
    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_WIDTH_EMU)
    presentation.slide_height = Emu(SLIDE_HEIGHT_EMU)
    try:
        add_svg_slide(presentation, extract_and_validate_svg(svg_markup))
    except Exception as exc:
        return {
            "status": "fail",
            "code": "EXPORT_SHAPES_PREFLIGHT",
            "detail": str(exc),
            "violations": [
                {
                    "code": "EXPORT_SHAPES_PREFLIGHT",
                    "detail": str(exc),
                }
            ],
        }
    slide = presentation.slides[-1]
    has_native_text = False
    has_picture = False
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
            has_native_text = True
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            has_picture = True
    return {
        "status": "pass",
        "code": "EXPORT_SHAPES_PREFLIGHT",
        "shape_count": len(slide.shapes),
        "has_native_text": has_native_text,
        "has_picture": has_picture,
        "violations": [],
    }
