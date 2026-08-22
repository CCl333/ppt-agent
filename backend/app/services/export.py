from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from posixpath import normpath

from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from app.services.svg import canvas_size, extract_and_validate_svg
from app.services.svg_pptx import SLIDE_HEIGHT_EMU, SLIDE_WIDTH_EMU, add_svg_slide

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
ASVG_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
IMAGE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def rasterize_svg_to_png(svg_markup: str) -> bytes:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("未安装 pymupdf，无法将 SVG 光栅化为 PNG") from exc

    try:
        document = fitz.open(stream=svg_markup.encode("utf-8"), filetype="svg")
    except Exception as exc:
        raise RuntimeError(f"SVG 光栅化失败: {exc}") from exc
    try:
        if document.page_count < 1:
            raise RuntimeError("SVG 光栅化失败: 没有可渲染页面")
        page = document[0]
        rect = page.rect
        if rect.width <= 0 or rect.height <= 0:
            raise RuntimeError("SVG 光栅化失败: 画布尺寸无效")
        zoom = 1920 / rect.width
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True)
        png = pixmap.tobytes("png")
    finally:
        document.close()
    if not png or not png.startswith(b"\x89PNG") or len(png) < 32:
        raise RuntimeError("SVG 光栅化失败: 输出不是有效 PNG")
    return png


def build_pptx(
    slides: list[tuple[str, str]],
    export_path: Path,
    *,
    mode: str = "shapes",
) -> Path:
    if not slides:
        raise RuntimeError("当前项目没有可导出的设计稿页面")
    if mode == "image":
        return _build_pptx_image(slides, export_path)
    if mode != "shapes":
        raise RuntimeError(f"不支持的导出模式: {mode}")
    return _build_pptx_shapes(slides, export_path)


def verify_pptx_package(path: Path, *, render_mode: str, expected_slide_count: int) -> dict[str, Any]:
    issues: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "ppt/presentation.xml" not in names:
            issues.append("缺少 ppt/presentation.xml")
        else:
            presentation = archive.read("ppt/presentation.xml").decode("utf-8")
            if f'cx="{SLIDE_WIDTH_EMU}"' not in presentation:
                issues.append("画布宽度与 16:9 约定不一致")
        slides = sorted(
            [
                name
                for name in names
                if name.startswith("ppt/slides/slide") and name.endswith(".xml") and "/_rels/" not in name
            ],
            key=lambda name: int(re.search(r"\d+", name).group(0) if re.search(r"\d+", name) else 0),
        )
        if len(slides) != expected_slide_count:
            issues.append(f"slide 数 {len(slides)} 与快照 {expected_slide_count} 不一致")
        joined = ""
        for slide_name in slides:
            joined += archive.read(slide_name).decode("utf-8")
            rels_name = f"ppt/slides/_rels/{Path(slide_name).name}.rels"
            if rels_name not in names:
                continue
            rels = archive.read(rels_name).decode("utf-8")
            for target in re.findall(r'Target="([^"]+)"', rels):
                if target.startswith("http"):
                    continue
                resolved = normpath(f"ppt/slides/{target}")
                if resolved not in names:
                    issues.append(f"幻灯片关系指向缺失对象: {target}")
        if render_mode == "shapes" and "svgBlip" in joined:
            issues.append("shapes 模式包含整页 SVG blip")
        if render_mode == "image" and "svgBlip" not in joined:
            issues.append("image 模式缺少 SVG blip")
        if render_mode == "image" and not any(name.startswith("ppt/media/") and name.endswith(".png") for name in names):
            issues.append("image 模式缺少 PNG 主图")
    if issues:
        raise RuntimeError("；".join(issues))
    return {"slide_count": len(slides), "issues": []}


def _build_pptx_shapes(slides: list[tuple[str, str]], export_path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_WIDTH_EMU)
    presentation.slide_height = Emu(SLIDE_HEIGHT_EMU)
    for page_code, svg_markup in slides:
        try:
            add_svg_slide(presentation, extract_and_validate_svg(svg_markup))
        except Exception as exc:
            raise RuntimeError(f"页面 {page_code} native shapes 导出失败: {exc}") from exc
    export_path = Path(export_path)
    export_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(export_path))
    return export_path


def _build_pptx_image(slides: list[tuple[str, str]], export_path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Emu(SLIDE_WIDTH_EMU)
    presentation.slide_height = Emu(SLIDE_HEIGHT_EMU)
    blank_layout = presentation.slide_layouts[6]
    svg_blobs: list[bytes] = []

    for _page_code, svg_markup in slides:
        svg = extract_and_validate_svg(svg_markup)
        png = rasterize_svg_to_png(svg)
        slide = presentation.slides.add_slide(blank_layout)
        left, top, width, height = _fit_picture_to_slide(
            slide_width_emu=int(presentation.slide_width),
            slide_height_emu=int(presentation.slide_height),
            svg_markup=svg,
        )
        slide.shapes.add_picture(io.BytesIO(png), left, top, width, height)
        svg_blobs.append(svg.encode("utf-8"))

    export_path = Path(export_path)
    export_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(export_path))
    inject_svg_blips(export_path, svg_blobs)
    return export_path


def inject_svg_blips(pptx_path: Path, svg_blobs: list[bytes]) -> None:
    source = pptx_path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(source), "r") as zin:
        files = {info.filename: zin.read(info.filename) for info in zin.infolist()}

    files["[Content_Types].xml"] = _ensure_svg_content_type(files["[Content_Types].xml"])
    slide_names = sorted(
        [name for name in files if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)],
        key=lambda name: int(re.search(r"\d+", name).group(0)),
    )
    if len(slide_names) != len(svg_blobs):
        raise RuntimeError(f"PPTX 页数 {len(slide_names)} 与 SVG 数 {len(svg_blobs)} 不一致")

    used_media = set()
    for name in files:
        match = re.search(r"ppt/media/image(\d+)", name)
        if match:
            used_media.add(int(match.group(1)))
    next_media = (max(used_media) if used_media else 0) + 1

    for slide_name, svg_bytes in zip(slide_names, svg_blobs):
        media_name = f"ppt/media/image{next_media}.svg"
        next_media += 1
        files[media_name] = svg_bytes
        rels_name = f"ppt/slides/_rels/{Path(slide_name).name}.rels"
        files[rels_name], svg_rid = _add_image_relationship(
            files[rels_name],
            f"../media/{Path(media_name).name}",
        )
        files[slide_name] = _attach_svg_blip(files[slide_name], svg_rid)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    pptx_path.write_bytes(buffer.getvalue())


def _fit_picture_to_slide(*, slide_width_emu: int, slide_height_emu: int, svg_markup: str) -> tuple[int, int, int, int]:
    width_px, height_px = canvas_size(svg_markup)
    if not width_px or not height_px:
        return 0, 0, slide_width_emu, slide_height_emu
    scale = min(slide_width_emu / width_px, slide_height_emu / height_px)
    width = int(width_px * scale)
    height = int(height_px * scale)
    left = int((slide_width_emu - width) / 2)
    top = int((slide_height_emu - height) / 2)
    return Emu(left), Emu(top), Emu(width), Emu(height)


def _ensure_svg_content_type(xml_bytes: bytes) -> bytes:
    root = etree.fromstring(xml_bytes)
    for default in root.findall(f"{{{CT_NS}}}Default"):
        if default.get("Extension") == "svg":
            return xml_bytes
    elem = etree.SubElement(root, f"{{{CT_NS}}}Default")
    elem.set("Extension", "svg")
    elem.set("ContentType", "image/svg+xml")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8")


def _add_image_relationship(xml_bytes: bytes, target: str) -> tuple[bytes, str]:
    root = etree.fromstring(xml_bytes)
    used: list[int] = []
    for rel in root.findall(f"{{{PKG_REL_NS}}}Relationship"):
        match = re.fullmatch(r"rId(\d+)", rel.get("Id") or "")
        if match:
            used.append(int(match.group(1)))
    rid = f"rId{(max(used) if used else 0) + 1}"
    rel = etree.SubElement(root, f"{{{PKG_REL_NS}}}Relationship")
    rel.set("Id", rid)
    rel.set("Type", IMAGE_REL_TYPE)
    rel.set("Target", target)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8"), rid


def _attach_svg_blip(xml_bytes: bytes, svg_rid: str) -> bytes:
    root = etree.fromstring(xml_bytes)
    blips = root.xpath("//a:blip", namespaces={"a": A_NS})
    if not blips:
        raise RuntimeError("幻灯片缺少 PNG 主图，无法写入 svgBlip")
    for blip in blips:
        if blip.xpath(".//asvg:svgBlip", namespaces={"asvg": ASVG_NS}):
            continue
        ext_lst = blip.find(f"{{{A_NS}}}extLst")
        if ext_lst is None:
            ext_lst = etree.SubElement(blip, f"{{{A_NS}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{A_NS}}}ext")
        ext.set("uri", ASVG_URI)
        svg_blip = etree.SubElement(ext, f"{{{ASVG_NS}}}svgBlip")
        svg_blip.set(f"{{{R_NS}}}embed", svg_rid)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8")
